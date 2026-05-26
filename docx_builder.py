import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def build_ebook_docx(
    title: str, content: str, cover_path: str, output_path: str
) -> None:
    doc = Document()
    _setup_page(doc)

    if cover_path and os.path.isfile(cover_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(cover_path, width=Inches(6))
        doc.add_page_break()

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title)
    run.font.size = Pt(28)
    run.font.bold = True
    doc.add_page_break()

    _add_markdown_content(doc, content)
    doc.save(output_path)


def build_value_enhancer_docx(data: dict, output_path: str) -> None:
    doc = Document()
    _setup_page(doc)

    h = doc.add_heading("Bonus Vault & Value Enhancer", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Bonus Bundle", level=1)
    for bonus in data.get("bonuses", []):
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{bonus.get('name', 'N/A')} ({bonus.get('format', 'N/A')})").bold = True
        p.add_run(f" – {bonus.get('benefit', 'N/A')}")

    doc.add_heading("Workbook Outline", level=1)
    for item in data.get("workbookOutline", []):
        p = doc.add_paragraph(style="List Number")
        p.add_run(item.get("section", "N/A")).bold = True
        p.add_run(f": {item.get('description', 'N/A')}")

    oto = data.get("oto", {})
    doc.add_heading(
        f"One-Time Offer — {oto.get('name', 'N/A')} (${oto.get('price', 'N/A')})",
        level=1,
    )
    doc.add_paragraph(oto.get("description", ""))
    for deliverable in oto.get("deliverables", []):
        doc.add_paragraph(str(deliverable), style="List Bullet")

    doc.save(output_path)


def _setup_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)


def _add_markdown_content(doc: Document, text: str) -> None:
    for line in text.splitlines():
        s = line.rstrip()
        if s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith(("* ", "- ")):
            _inline(doc.add_paragraph(style="List Bullet"), s[2:])
        elif re.match(r"^\d+\.\s", s):
            _inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s", "", s))
        elif s:
            _inline(doc.add_paragraph(), s)


def _inline(paragraph, text: str) -> None:
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            paragraph.add_run(chunk[1:-1]).italic = True
        else:
            paragraph.add_run(chunk)
